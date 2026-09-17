"""Generic CSV/Excel bulk-import runner shared by resources that need it
(employees, leave records). Each data row is validated and inserted in its
own SAVEPOINT so one bad row doesn't abort rows already committed in the
same request.
"""
import csv
import io
from typing import Callable, TypeVar

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.schemas import BulkImportError, BulkImportResult

T = TypeVar("T", bound=BaseModel)

# SECURITY: cap the upload size read into memory. Without this, an
# arbitrarily large file (accidental or malicious — this endpoint is
# admin-only, but "trusted role" still isn't "trust the file size") would
# be fully buffered as bytes and then again as a decoded string, which is an
# easy memory-exhaustion DoS vector. 10 MB is generous for a CSV/Excel of
# HR/leave records (tens of thousands of rows) while ruling out
# multi-hundred-MB uploads. See SECURITY_REPORT.md.
MAX_BULK_IMPORT_BYTES = 10 * 1024 * 1024


def _clean(value):
    return value.strip() if isinstance(value, str) else value


def _parse_csv_rows(raw_bytes: bytes):
    raw = raw_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(raw))
    for i, row in enumerate(reader, start=2):
        yield i, {k: _clean(v) for k, v in row.items() if k and v not in (None, "")}


def _parse_xlsx_rows(raw_bytes: bytes):
    wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return
    for i, raw_row in enumerate(rows, start=2):
        row = {headers[j]: _clean(v) for j, v in enumerate(raw_row) if j < len(headers) and headers[j] and v not in (None, "")}
        if row:
            yield i, row
    wb.close()


def _is_xlsx(file: UploadFile) -> bool:
    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        return True
    return file.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def run_csv_bulk_import(
    db: Session,
    file: UploadFile,
    schema: type[T],
    insert_row: Callable[[Session, T], None],
) -> BulkImportResult:
    raw_bytes = await file.read(MAX_BULK_IMPORT_BYTES + 1)
    if len(raw_bytes) > MAX_BULK_IMPORT_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File too large (max {MAX_BULK_IMPORT_BYTES // (1024 * 1024)} MB)",
        )

    if _is_xlsx(file):
        try:
            rows = list(_parse_xlsx_rows(raw_bytes))
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Could not read Excel file: {exc}")
    else:
        try:
            rows = list(_parse_csv_rows(raw_bytes))
        except UnicodeDecodeError:
            # A file that isn't UTF-8 text is a caller error at a system
            # boundary, not a server fault. Mirrors the .xlsx branch above:
            # without this, UnicodeDecodeError escapes as a 500.
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Could not read CSV file: it is not valid UTF-8 text. Save it as UTF-8 (or upload an .xlsx file) and try again.",
            )
        except csv.Error as exc:
            # Decodes as text but the csv module refuses to parse it — most
            # reachably a field longer than csv.field_size_limit() (131072
            # chars), which the 10 MB upload cap above happily allows
            # through. Same boundary reasoning as the decode guard.
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Could not read CSV file: {exc}")

    errors: list[BulkImportError] = []
    created = 0
    for i, cleaned in rows:
        try:
            payload = schema(**cleaned)
        except ValidationError as exc:
            errors.append(BulkImportError(row=i, message=exc.errors()[0]["msg"]))
            continue

        try:
            with db.begin_nested():
                insert_row(db, payload)
        except ValueError as exc:
            # App-raised, already a clean/safe message (e.g. "Location not
            # found") — fine to return to the (admin-only) caller as-is.
            errors.append(BulkImportError(row=i, message=str(exc)))
            continue
        except IntegrityError:
            # SECURITY: do not surface exc.orig — that's the raw driver/DB
            # error text (constraint/index names, table names, sometimes
            # verbatim column values), which is internal implementation
            # detail we shouldn't echo back over the API even to a trusted
            # role. A clean, generic message is enough for HR to know which
            # row to fix. See SECURITY_REPORT.md.
            errors.append(
                BulkImportError(
                    row=i,
                    message="Row conflicts with an existing record or references a value that doesn't exist",
                )
            )
            continue
        created += 1

    db.commit()
    return BulkImportResult(created=created, errors=errors)
