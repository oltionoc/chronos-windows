"""Payroll Excel exports — run summary (one row per employee) and individual
payslip, both generated on demand, not persisted to disk."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models import PayrollAdjustment, PayrollRun, PayrollRunLine

_HEADERS = [
    "Employee",
    "Base Salary (EUR)",
    "Late (min)",
    "Lateness Penalty (EUR)",
    "Overtime (min)",
    "Overtime Bonus (EUR)",
    "Absence Days",
    "Absence Deduction (EUR)",
    "Paid Leave (days)",
    "Unpaid Leave (days)",
    "Net Pay (EUR)",
]


def generate_payroll_run_xlsx(run: PayrollRun, lines: list[PayrollRunLine]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{run.period_year}-{run.period_month:02d}"

    ws.append(_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for l in lines:
        ws.append(
            [
                f"{l.employee.first_name} {l.employee.last_name}" if l.employee else "",
                float(l.base_salary_eur),
                l.total_late_minutes,
                float(l.total_lateness_penalty_eur),
                l.total_overtime_minutes,
                float(l.total_overtime_bonus_eur),
                float(l.total_absence_days),
                float(l.total_absence_deduction_eur),
                float(l.paid_leave_days),
                float(l.unpaid_leave_days),
                float(l.net_pay_eur),
            ]
        )

    total_row = len(lines) + 2
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for col in (2, 4, 6, 8, 11):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{total_row - 1})")
        cell.font = Font(bold=True)

    for col in range(1, len(_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_payslip_xlsx(line: PayrollRunLine, adjustments: list[PayrollAdjustment]) -> bytes:
    emp = line.employee
    run = line.run
    wb = Workbook()
    ws = wb.active
    ws.title = "Payslip"

    rows = [
        ("Employee", f"{emp.first_name} {emp.last_name}" if emp else ""),
        ("Employee Code", emp.employee_code if emp else ""),
        ("Location", run.location.name if run.location else ""),
        ("Period", f"{run.period_year}-{run.period_month:02d}"),
        ("", ""),
        ("Base Salary (EUR)", float(line.base_salary_eur)),
        ("Late (min)", line.total_late_minutes),
        ("Lateness Penalty (EUR)", float(line.total_lateness_penalty_eur)),
        ("Overtime (min)", line.total_overtime_minutes),
        ("Overtime Bonus (EUR)", float(line.total_overtime_bonus_eur)),
        ("Absence Days", float(line.total_absence_days)),
        ("Absence Deduction (EUR)", float(line.total_absence_deduction_eur)),
        ("Paid Leave (days)", float(line.paid_leave_days)),
        ("Unpaid Leave (days)", float(line.unpaid_leave_days)),
        ("Net Pay (EUR)", float(line.net_pay_eur)),
    ]
    for label, value in rows:
        ws.append([label, value])
    for row in ws.iter_rows(min_row=1, max_row=len(rows)):
        row[0].font = Font(bold=True)

    if adjustments:
        ws.append([])
        ws.append(["Adjustments"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append(["Type", "Reason", "Amount (EUR)"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for a in adjustments:
            ws.append([a.type, a.reason, float(a.amount_eur)])

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
